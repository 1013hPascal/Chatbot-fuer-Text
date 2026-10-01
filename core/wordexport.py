"""Textentwurf als Word-Dokument (.docx) speichern."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def safe_filename(name: str, default: str = "Text") -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip(" .")
    name = re.sub(r"\.docx$", "", name, flags=re.I)
    return (name or default)[:80]


def save_docx(path: str | Path, title: str, text: str) -> Path:
    """Absätze durch Leerzeilen, Zeilenumbrüche innerhalb eines Absatzes bleiben erhalten
    (Adressen, Gedichte), Zeilen mit '- ' werden zu Aufzählungspunkten."""
    doc = Document()
    doc.core_properties.title = title
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = block.split("\n")
        if all(line.lstrip().startswith(("- ", "• ")) for line in lines):
            for line in lines:
                doc.add_paragraph(line.lstrip()[2:].strip(), style="List Bullet")
            continue
        paragraph = doc.add_paragraph()
        for i, line in enumerate(lines):
            run = paragraph.add_run(line.rstrip())
            if i < len(lines) - 1:
                run.add_break()
    target = Path(path)
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    doc.save(str(target))
    return target
