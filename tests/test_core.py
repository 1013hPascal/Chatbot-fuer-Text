"""Tests für Modell-Ausgabe, Werkzeug-Schleife, Ordnerzugriff, Anhänge, Word-Export und Speicher."""
import threading

import pytest
from docx import Document

from core import prompts, websearch
from core.config import Config
from core.conversation import SHORTENED, ChatSession
from core.files import AttachedFile, FileReadError, read_file
from core.folders import FolderAccess, FolderError
from core.llm import Cancelled, LLMError, ToolCall, create_backend
from core.store import Store, title_from_question
from core.tools import ToolBox
from core.wordexport import safe_filename, save_docx
from tests.conftest import FakeBackend, content, tool_call


# -- Auswertung der Modell-Ausgabe --------------------------------------------------------
def test_parse_reply_extracts_document():
    raw = ('Ich habe das Anschreiben erstellt.\n<dokument titel="Bewerbung">\nSehr geehrte '
           'Damen und Herren,\n\nhiermit bewerbe ich mich.\n</dokument>')
    reply = prompts.parse_reply(raw)
    assert reply.answer == "Ich habe das Anschreiben erstellt."
    assert reply.doc_title == "Bewerbung"
    assert reply.doc_text.startswith("Sehr geehrte Damen und Herren,")
    assert reply.word_filename is None


def test_parse_reply_without_document_and_markdown_is_cleaned():
    reply = prompts.parse_reply("**Wichtig:** Das ist `Code` und\n\n\n\n## Titel\n* Punkt")
    assert reply.doc_text is None
    assert "*" not in reply.answer and "`" not in reply.answer and "#" not in reply.answer
    assert "- Punkt" in reply.answer


def test_parse_reply_unclosed_block_and_think_tags():
    raw = '<think>Ich überlege lange.</think>Fertig.\n<dokument titel="Brief">Text ohne Ende'
    reply = prompts.parse_reply(raw)
    assert reply.answer == "Fertig."
    assert reply.doc_text == "Text ohne Ende"
    assert "überlege" not in reply.answer


def test_parse_reply_word_offer():
    reply = prompts.parse_reply('Bitte sehr.\n<dokument titel="Brief">Hallo</dokument>\n'
                                '<word_dokument dateiname="Bewerbung"/>')
    assert reply.word_filename == "Bewerbung"
    reply = prompts.parse_reply('Bitte.\n<dokument titel="Mein Brief">Hallo</dokument>'
                                '<word_dokument/>')
    assert reply.word_filename == "Mein Brief"


def test_multiple_document_blocks_keep_the_last():
    raw = '<dokument titel="A">alt</dokument>\nText\n<dokument titel="B">neu</dokument>'
    reply = prompts.parse_reply(raw)
    assert reply.doc_text == "neu" and reply.doc_title == "B"
    assert "alt" not in reply.answer


def test_system_prompt_reflects_settings():
    short = prompts.build_system(False, False)
    long_web = prompts.build_system(True, True, [("a.txt", "INHALT")], ["Verträge"])
    assert prompts.LENGTH_SHORT in short and "keinen Zugriff auf das Internet" in short
    assert "ordner_auflisten" not in short
    assert prompts.LENGTH_LONG in long_web and "internet_suche" in long_web
    assert "a.txt" in long_web and "INHALT" in long_web
    assert "Verträge" in long_web and "ordner_auflisten" in long_web


# -- Unterhaltung -------------------------------------------------------------------------
@pytest.fixture
def session(cfg):
    cfg.web_search = False
    return ChatSession(cfg, FakeBackend(cfg))


def test_answer_and_document_flow(session):
    session.backend.script = [content('Erledigt.\n<dokument titel="Brief">Liebe Anna</dokument>')]
    result = session.ask("Schreibe einen Brief", "")
    assert result.answer == "Erledigt." and result.doc_text == "Liebe Anna"
    assert result.doc_title == "Brief" and not result.web_sources


def test_document_only_reply_gets_a_spoken_answer(session):
    session.backend.script = [content('<dokument titel="Brief">Liebe Anna</dokument>'),
                              content('<dokument>Lieber Anton</dokument>'),
                              content("<word_dokument/>")]
    assert session.ask("Brief", "").answer == prompts.DOC_CREATED
    assert session.ask("Ändere", "Liebe Anna").answer == prompts.DOC_CHANGED
    assert session.ask("Als Word", "Lieber Anton").answer == prompts.WORD_READY


def test_memory_holds_plain_answers_without_internal_notes(session):
    session.backend.script = [content('<dokument titel="B">Text</dokument>'), content("Ok.")]
    session.ask("Erste", "")
    session.ask("Zweite", "Text")
    assert session.backend.calls[1]["messages"][2]["content"] == prompts.DOC_CREATED


def test_current_text_and_history_are_sent_to_model(session):
    session.backend.script = [content("Erste Antwort."), content("Zweite Antwort.")]
    session.ask("Erste Frage", "")
    session.ask("Zweite Frage", text="MEIN ENTWURF")
    messages = session.backend.calls[1]["messages"]
    assert [m["role"] for m in messages] == ["system", "user", "assistant", "user"]
    assert messages[1]["content"] == "Erste Frage"
    assert "MEIN ENTWURF" in messages[3]["content"] and "Zweite Frage" in messages[3]["content"]
    assert "MEIN ENTWURF" not in session.backend.calls[0]["messages"][-1]["content"]


def test_load_restores_memory(session):
    session.load([("Alt-Frage", "Alt-Antwort")])
    session.ask("Neu", "")
    sent = session.backend.calls[0]["messages"]
    assert sent[1]["content"] == "Alt-Frage" and sent[2]["content"] == "Alt-Antwort"


def test_reset_forgets_everything(session, tmp_path):
    session.ask("Frage", "")
    session.add_attachment(AttachedFile("a.txt", "Inhalt"))
    session.folders.grant(tmp_path)
    session.reset()
    session.ask("Neu", "")
    assert len(session.backend.calls[-1]["messages"]) == 2
    assert "Inhalt" not in session.backend.calls[-1]["messages"][0]["content"]
    assert session.folders.roots == []


def test_settings_reach_the_model_call(cfg):
    cfg.web_search, cfg.thinking, cfg.long_answers = False, True, True
    backend = FakeBackend(cfg)
    ChatSession(cfg, backend).ask("Hallo", "")
    call = backend.calls[0]
    assert call["think"] is True and call["tools"] is None
    assert call["max_tokens"] == cfg.tokens_long + cfg.tokens_thinking_extra
    assert prompts.LENGTH_LONG in call["messages"][0]["content"]


def test_short_mode_without_thinking_has_small_budget(cfg):
    cfg.web_search = False
    backend = FakeBackend(cfg)
    ChatSession(cfg, backend).ask("Hallo", "")
    assert backend.calls[0]["think"] is False
    assert backend.calls[0]["max_tokens"] == cfg.tokens_normal


def test_web_search_tool_loop_and_sources(cfg):
    cfg.web_search, cfg.web_prefetch = True, False
    backend = FakeBackend(cfg, [
        tool_call("internet_suche", anfrage="Wetter Freiburg"),
        tool_call("webseite_lesen", url="https://wetter.example/freiburg"),
        content("Morgen wird es sonnig."),
    ])
    box = ToolBox(search_fn=lambda q: [{"title": "Wetter", "url": "https://wetter.example/freiburg",
                                        "snippet": "sonnig"}],
                  read_fn=lambda url: ("Wetter Freiburg", "Morgen sonnig, 24 Grad."))
    statuses = []
    result = ChatSession(cfg, backend, box).ask("Wie wird das Wetter?", "",
                                                on_status=statuses.append)
    assert result.web_sources == ["https://wetter.example/freiburg"]
    assert result.answer.endswith("Quellen aus dem Internet: https://wetter.example/freiburg")
    assert "Internetrecherche: Wetter Freiburg" in statuses
    assert "Seite wird gelesen: wetter.example" in statuses
    tool_messages = [m for m in backend.calls[2]["messages"] if m["role"] == "tool"]
    assert len(tool_messages) == 2 and "24 Grad" in tool_messages[1]["content"]
    assert [t["function"]["name"] for t in backend.calls[0]["tools"]] == ["internet_suche",
                                                                          "webseite_lesen"]


def test_search_hits_are_sources_when_no_page_was_read(cfg):
    backend = FakeBackend(cfg, [tool_call("internet_suche", anfrage="x"), content("Antwort.")])
    box = ToolBox(search_fn=lambda q: [{"title": "T", "url": f"https://a.example/{i}",
                                        "snippet": "s"} for i in range(5)])
    result = ChatSession(cfg, backend, box).ask("Frage", "")
    assert result.answer.count("https://") == 3


def test_tool_errors_are_given_to_model_not_raised(cfg):
    def failing(_):
        raise websearch.WebError("Keine Verbindung")
    backend = FakeBackend(cfg, [tool_call("internet_suche", anfrage="x"),
                                content("Ich konnte nicht suchen.")])
    result = ChatSession(cfg, backend, ToolBox(search_fn=failing)).ask("Frage", "")
    assert "Fehler: Keine Verbindung" in backend.calls[1]["messages"][-1]["content"]
    assert result.answer == "Ich konnte nicht suchen."


# -- Tavily mit DuckDuckGo als Rückfall ---------------------------------------------------------
HIT = {"title": "T", "url": "https://a.example/1", "snippet": "s"}


def _searcher(cfg, tavily, ddg=None):
    return websearch.Searcher(cfg, tavily_fn=tavily, ddg_fn=ddg or (lambda q: [{**HIT, "title": "DDG"}]))


def test_searcher_without_key_uses_duckduckgo(cfg, monkeypatch):
    monkeypatch.delenv("TAVILY_API_KEY", raising=False)
    cfg.tavily_key = ""
    searcher = _searcher(cfg, lambda *a, **k: pytest.fail("Tavily ohne Schlüssel"))
    assert searcher("x")[0]["title"] == "DDG" and searcher.take_notices() == []


def test_searcher_with_key_uses_tavily_and_passes_key_and_depth(cfg):
    cfg.tavily_key, cfg.tavily_depth = " tvly-1 ", "advanced"
    seen = {}

    def tavily(query, key, depth):
        seen.update(query=query, key=key, depth=depth)
        return [HIT]
    assert _searcher(cfg, tavily,
                     lambda q: pytest.fail("kein Rückfall nötig"))("Wetter")[0]["title"] == "T"
    assert seen == {"query": "Wetter", "key": "tvly-1", "depth": "advanced"}


def test_searcher_reads_key_from_environment_and_ignores_bad_depth(cfg, monkeypatch):
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-env")
    cfg.tavily_key, cfg.tavily_depth = "", "unsinn"
    seen = {}
    _searcher(cfg, lambda q, key, depth: seen.update(key=key, depth=depth) or [HIT])("x")
    assert seen == {"key": "tvly-env", "depth": "basic"}


def test_searcher_falls_back_and_reports_reason_once(cfg):
    cfg.tavily_key = "tvly-1"

    def broken(*a, **k):
        raise websearch.TavilyError("Das Tavily-Kontingent ist aufgebraucht")
    searcher = _searcher(cfg, broken)
    assert searcher("x")[0]["title"] == "DDG"
    assert searcher.take_notices() == ["Das Tavily-Kontingent ist aufgebraucht. "
                                       "Es wird DuckDuckGo genutzt."]
    searcher("y")
    assert searcher.take_notices() == []          # derselbe Grund wird nur einmal gemeldet


def test_searcher_uses_duckduckgo_when_tavily_finds_nothing(cfg):
    cfg.tavily_key = "tvly-1"
    searcher = _searcher(cfg, lambda *a, **k: [])
    assert searcher("x")[0]["title"] == "DDG" and searcher.take_notices() == []


class _Response:
    def __init__(self, status=200, data=None):
        self.status_code, self._data = status, data

    def json(self):
        if self._data is None:
            raise ValueError("kein JSON")
        return self._data


def test_search_tavily_parses_results_and_sends_bearer_key(monkeypatch):
    sent = {}

    def post(url, json, headers, timeout):
        sent.update(url=url, json=json, headers=headers)
        return _Response(200, {"results": [{"title": "A", "url": "https://a.example", "content": "Text",
                                            "score": 0.9}, {"title": "ohne Adresse"}]})
    monkeypatch.setattr(websearch.httpx, "post", post)
    hits = websearch.search_tavily("Frage", "tvly-1", max_results=3, depth="advanced")
    assert hits == [{"title": "A", "url": "https://a.example", "snippet": "Text"}]
    assert sent["headers"] == {"Authorization": "Bearer tvly-1"}
    assert sent["json"]["query"] == "Frage" and sent["json"]["search_depth"] == "advanced"
    assert sent["json"]["max_results"] == 3 and sent["url"] == websearch.TAVILY_URL


@pytest.mark.parametrize("status,word", [(401, "abgelehnt"), (429, "Kontingent"),
                                         (432, "Kontingent"), (433, "Kontingent"),
                                         (500, "Fehler 500")])
def test_search_tavily_errors_become_tavily_error(monkeypatch, status, word):
    monkeypatch.setattr(websearch.httpx, "post", lambda *a, **k: _Response(status, {}))
    with pytest.raises(websearch.TavilyError, match=word):
        websearch.search_tavily("x", "k")


def test_search_tavily_network_and_garbage_answers(monkeypatch):
    def offline(*a, **k):
        raise websearch.httpx.ConnectError("weg")
    monkeypatch.setattr(websearch.httpx, "post", offline)
    with pytest.raises(websearch.TavilyError, match="nicht erreichbar"):
        websearch.search_tavily("x", "k")
    monkeypatch.setattr(websearch.httpx, "post", lambda *a, **k: _Response(200, None))
    with pytest.raises(websearch.TavilyError, match="unlesbar"):
        websearch.search_tavily("x", "k")


def test_search_notice_is_announced_during_the_turn(cfg):
    cfg.tavily_key = "tvly-1"

    def broken(*a, **k):
        raise websearch.TavilyError("Der Tavily-Schlüssel wurde abgelehnt")
    box = ToolBox(search_fn=_searcher(cfg, broken))
    backend = FakeBackend(cfg, [tool_call("internet_suche", anfrage="x"), content("Antwort.")])
    statuses = []
    ChatSession(cfg, backend, box).ask("Frage", "", on_status=statuses.append)
    assert "Der Tavily-Schlüssel wurde abgelehnt. Es wird DuckDuckGo genutzt." in statuses


def test_default_session_searches_with_the_config_key(cfg):
    session = ChatSession(cfg, FakeBackend(cfg, []))
    assert isinstance(session.toolbox._search, websearch.Searcher)
    assert session.toolbox._search.cfg is cfg


# -- Modellfähigkeiten (Ollama) -------------------------------------------------------------
def _ollama_with(cfg, handler):
    import httpx
    from core.ollama_backend import OllamaBackend
    backend = OllamaBackend(cfg)
    backend._http = httpx.Client(base_url="http://test", transport=httpx.MockTransport(handler))
    return backend


def test_ollama_capabilities_are_read_once_per_model_and_limit_the_request(cfg):
    import httpx
    import json
    seen = {"show": 0, "chat": []}

    def handler(request):
        body = json.loads(request.content)
        if request.url.path == "/api/show":
            seen["show"] += 1
            return httpx.Response(200, json={"capabilities": ["completion", "vision"]})
        seen["chat"].append(body)
        return httpx.Response(200, content=b'{"message": {"content": "ok"}, "done": true}\n')
    cfg.chat_model = "qwen2.5vl:7b"
    backend = _ollama_with(cfg, handler)
    assert backend.capabilities() == {"completion", "vision"}
    chunks = list(backend.chat_stream([{"role": "user", "content": "x"}],
                                      [{"type": "function"}], think=True))
    assert [c.text for c in chunks] == ["ok"]
    assert "tools" not in seen["chat"][0] and "think" not in seen["chat"][0]   # nicht unterstützt
    assert seen["show"] == 1                                                      # nur einmal gefragt
    cfg.chat_model = "gemma4:12b"
    backend.capabilities()
    assert seen["show"] == 2                                                      # je Modell einmal


def test_ollama_capabilities_unknown_when_show_fails_and_unload(cfg):
    import httpx
    calls = []

    def handler(request):
        calls.append((request.url.path, request.content))
        return httpx.Response(500)
    backend = _ollama_with(cfg, handler)
    assert backend.capabilities() == {"tools", "thinking", "vision"}   # nichts sperren, nichts merken
    assert backend._caps == {}
    backend.unload("altes:1b")
    assert calls[-1][0] == "/api/generate" and b'"keep_alive": 0' in calls[-1][1].replace(b":0", b": 0")


def test_tool_loop_is_limited(cfg):
    cfg.max_tool_steps = 2
    loop = [tool_call("internet_suche", anfrage="x")] * 5 + [content("Ende.")]
    backend = FakeBackend(cfg, loop)
    ChatSession(cfg, backend, ToolBox(search_fn=lambda q: [])).ask("Frage", "")
    assert len(backend.calls) == 3
    assert backend.calls[-1]["tools"] is None          # letzter Schritt ohne Werkzeuge


def test_no_tools_offered_when_web_off_and_no_folder(cfg):
    cfg.web_search = False
    backend = FakeBackend(cfg)
    ChatSession(cfg, backend).ask("Frage", "")
    assert backend.calls[0]["tools"] is None


def test_empty_output_gives_helpful_message(session):
    session.backend.script = [[]]
    assert session.ask("Frage", "").answer == prompts.NO_ANSWER


def test_cancel_raises(cfg):
    cfg.web_search = False
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(Cancelled):
        ChatSession(cfg, FakeBackend(cfg, delay=1.0)).ask("Frage", "", cancel=cancel)


def test_old_turns_are_dropped_when_context_is_full(cfg):
    cfg.web_search = False
    cfg.num_ctx = 4096
    backend = FakeBackend(cfg)
    session = ChatSession(cfg, backend)
    session.turns = [(f"Frage {i} " + "x" * 800, "Antwort " + "y" * 800) for i in range(30)]
    session.ask("Neu", "")
    sent = backend.calls[0]["messages"]
    assert 2 < len(sent) < 62
    assert "Frage 29" in sent[-3]["content"]           # neueste Runde bleibt erhalten


def test_attachments_are_included_and_limited(cfg):
    cfg.web_search = False
    cfg.max_attachment_chars = 100
    backend = FakeBackend(cfg)
    session = ChatSession(cfg, backend)
    assert not session.add_attachment(AttachedFile("klein.txt", "kurzer Text"))
    assert session.add_attachment(AttachedFile("gross.txt", "wort " * 100))      # gekürzt
    assert session.attachment_chars() <= 100 + 20
    session.ask("Frage", "")
    system = backend.calls[0]["messages"][0]["content"]
    assert "klein.txt" in system and "kurzer Text" in system and "gross.txt" in system
    session.remove_attachment(0)
    assert [a.name for a in session.attachments] == ["gross.txt"]


# -- Freigegebene Ordner -----------------------------------------------------------------
@pytest.fixture
def folder(tmp_path):
    root = tmp_path / "Unterlagen"
    (root / "Verträge").mkdir(parents=True)
    (root / "Verträge" / "Miete.txt").write_text("Die Kaution beträgt 2550 Euro.", encoding="utf-8")
    (root / "notiz.txt").write_text("Termin am Montag um 10 Uhr.", encoding="utf-8")
    (root / "lang.txt").write_text("abcdefghij" * 1500, encoding="utf-8")
    (root / ".versteckt").mkdir()
    (root / ".versteckt" / "geheim.txt").write_text("nicht zeigen", encoding="utf-8")
    (root / "bild.png").write_bytes(b"\x89PNG")
    save_docx(root / "Brief.docx", "Brief", "Sehr geehrte Frau Weiß.")
    (tmp_path / "außerhalb.txt").write_text("GEHEIM", encoding="utf-8")
    access = FolderAccess()
    access.grant(root)
    return access


def test_folder_lists_only_supported_visible_files(folder):
    text = folder.list_text()
    assert "Verträge/Miete.txt" in text and "notiz.txt" in text and "Brief.docx" in text
    assert "geheim" not in text and "bild.png" not in text


def test_folder_reads_by_relative_path_and_by_name(folder):
    assert "2550 Euro" in folder.read("Verträge/Miete.txt")[1]
    assert "2550 Euro" in folder.read("Miete.txt")[1]
    assert "Frau Weiß" in folder.read("Brief.docx")[1]


def test_folder_reads_long_files_in_chunks(folder):
    label, first = folder.read("lang.txt")
    assert "Zeichen 0 bis 8000 von 15000" in first and "start=8000" in first
    _, second = folder.read("lang.txt", start=8000)
    assert "Zeichen 8000 bis 15000 von 15000" not in second and "start=" not in second


@pytest.mark.parametrize("name", ["../außerhalb.txt", "..\\außerhalb.txt", ".versteckt/geheim.txt",
                                  "", "gibtesnicht.txt"])
def test_folder_refuses_paths_outside_or_missing(folder, tmp_path, name):
    try:
        text = folder.read(name)[1]
    except FolderError:
        return
    assert "GEHEIM" not in text and "nicht zeigen" not in text
    pytest.fail("hätte abgelehnt werden müssen")


def test_folder_refuses_absolute_path_outside(folder, tmp_path):
    with pytest.raises(FolderError):
        folder.read(str(tmp_path / "außerhalb.txt"))


def test_folder_ambiguous_name(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    (tmp_path / "a" / "x.txt").write_text("eins", encoding="utf-8")
    (tmp_path / "b" / "x.txt").write_text("zwei", encoding="utf-8")
    access = FolderAccess()
    access.grant(tmp_path)
    with pytest.raises(FolderError, match="nicht eindeutig"):
        access.read("x.txt")
    assert "eins" in access.read("a/x.txt")[1]


def test_folder_search(folder):
    assert "Miete.txt" in folder.search("Kaution") and "2550" in folder.search("kaution")
    assert "Kein Treffer" in folder.search("Zeppelin")
    assert "Termin am Montag" in folder.search("montag termin")


def test_grant_of_missing_folder_fails(tmp_path):
    with pytest.raises(FolderError):
        FolderAccess().grant(tmp_path / "gibtesnicht")


def test_model_uses_folder_tools_and_files_are_listed_as_sources(cfg, folder):
    cfg.web_search = False
    backend = FakeBackend(cfg, [tool_call("in_dateien_suchen", begriff="Kaution"),
                                tool_call("datei_lesen", datei="Verträge/Miete.txt"),
                                content("Die Kaution beträgt 2550 Euro.")])
    session = ChatSession(cfg, backend, ToolBox(folders=folder))
    result = session.ask("Wie hoch ist die Kaution?", "")
    assert [t["function"]["name"] for t in backend.calls[0]["tools"]] == \
        ["ordner_auflisten", "in_dateien_suchen", "datei_lesen"]
    assert "Unterlagen" in backend.calls[0]["messages"][0]["content"]
    assert result.file_sources == ["Verträge/Miete.txt"]
    assert result.answer.endswith("Gelesene Dateien: Verträge/Miete.txt")
    last_tool = [m for m in backend.calls[2]["messages"] if m["role"] == "tool"][-1]
    assert "2550 Euro" in last_tool["content"]


def test_files_with_search_hits_count_as_sources_when_none_was_read(cfg, folder):
    cfg.web_search = False
    backend = FakeBackend(cfg, [tool_call("in_dateien_suchen", begriff="Kaution"),
                                content("Die Kaution beträgt 2550 Euro.")])
    result = ChatSession(cfg, backend, ToolBox(folders=folder)).ask("Kaution?", "")
    assert result.file_sources == ["Verträge/Miete.txt"]
    assert result.answer.endswith("Gelesene Dateien: Verträge/Miete.txt")


def test_folder_tool_errors_go_to_the_model(cfg, folder):
    cfg.web_search = False
    backend = FakeBackend(cfg, [tool_call("datei_lesen", datei="../außerhalb.txt"),
                                content("Nicht möglich.")])
    ChatSession(cfg, backend, ToolBox(folders=folder)).ask("Lies sie", "")
    result_text = backend.calls[1]["messages"][-1]["content"]
    assert result_text.startswith("Fehler:") and "GEHEIM" not in result_text


def test_old_tool_results_are_shortened_when_context_is_full(cfg, folder):
    cfg.web_search = False
    cfg.num_ctx = 4096
    cfg.tokens_normal = 500
    steps = [tool_call("datei_lesen", datei="lang.txt")] * 4 + [content("Fertig.")]
    backend = FakeBackend(cfg, steps)
    ChatSession(cfg, backend, ToolBox(folders=folder)).ask("Lies alles", "")
    last = backend.calls[-1]["messages"]
    tool_results = [m["content"] for m in last if m["role"] == "tool"]
    assert tool_results[-1] != SHORTENED and SHORTENED in tool_results


def test_toolbox_specs_depend_on_switches():
    box = ToolBox()
    assert box.specs(False, False) == []
    assert len(box.specs(True, False)) == 2 and len(box.specs(False, True)) == 3
    assert box.run(ToolCall("gibtesnicht")).startswith("Fehler")


# -- Dateien, Word, Web-Hilfen -------------------------------------------------------------
def test_read_txt_docx_and_unsupported(tmp_path):
    txt = tmp_path / "notiz.txt"
    txt.write_text("Grüße aus Freiburg", encoding="utf-8")
    assert read_file(txt).text == "Grüße aus Freiburg"
    old = tmp_path / "alt.txt"
    old.write_bytes("Größe".encode("cp1252"))
    assert read_file(old).text == "Größe"

    docx = tmp_path / "brief.docx"
    save_docx(docx, "Brief", "Erster Absatz.\n\nZweiter Absatz.")
    assert read_file(docx).text.splitlines()[0] == "Erster Absatz."

    with pytest.raises(FileReadError):
        read_file(tmp_path / "bild.png")
    empty = tmp_path / "leer.txt"
    empty.write_text("   ")
    with pytest.raises(FileReadError):
        read_file(empty)


def test_read_pdf(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Inhalt der PDF")
    path = tmp_path / "a.pdf"
    doc.save(path)
    assert "Inhalt der PDF" in read_file(path).text


def test_docx_roundtrip_keeps_paragraphs_lines_and_bullets(tmp_path):
    text = "Anna Müller\nHauptstraße 1\n\nSehr geehrte Damen und Herren,\n\n- Punkt eins\n- Punkt zwei"
    path = save_docx(tmp_path / "Brief", "Brief", text)
    assert path.suffix == ".docx" and path.exists()
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[0] == "Anna Müller\nHauptstraße 1"
    assert paragraphs[1] == "Sehr geehrte Damen und Herren,"
    assert [p.style.name for p in doc.paragraphs[2:]] == ["List Bullet", "List Bullet"]


def test_safe_filename():
    assert safe_filename('Brief: an "Anna"/Bob?.docx') == "Brief an AnnaBob"
    assert safe_filename("   ") == "Text"


def test_html_to_text_skips_scripts_and_navigation():
    html = ("<html><head><title>Titel</title><script>var x=1;</script></head><body><nav>Menü</nav>"
            "<h1>Überschrift</h1><p>Erster Absatz mit Inhalt.</p><footer>Impressum</footer></body></html>")
    title, text = websearch.html_to_text(html)
    assert title == "Titel"
    assert "Überschrift" in text and "Erster Absatz" in text
    assert "var x" not in text and "Menü" not in text and "Impressum" not in text


@pytest.mark.parametrize("url", ["file:///C:/Windows/win.ini", "ftp://example.com/a",
                                 "http://127.0.0.1:11434/api/tags", "http://localhost/x",
                                 "http://192.168.1.1/", "http://169.254.169.254/latest"])
def test_internal_and_non_http_addresses_are_refused(url):
    with pytest.raises(websearch.WebError):
        websearch.check_url(url)


# -- Speicher: Projekte, Chats, Verlauf -----------------------------------------------------
def test_chats_and_exchanges(store):
    chat = store.create_chat("Erste Frage")
    a = store.add_exchange(chat.id, "F1", "A1")
    b = store.add_exchange(chat.id, "F2", "A2", "Kommentar", True, "Brief", "Brief")
    entries = store.exchanges(chat.id)
    assert [e.id for e in entries] == [b.id, a.id]                    # neueste zuerst
    assert entries[0].is_doc and entries[0].comment == "Kommentar" and entries[0].word_filename == "Brief"
    store.update_exchange_answer(a.id, "A1 geändert")
    assert store.exchanges(chat.id)[1].answer == "A1 geändert"


def test_chat_order_rename_and_cascade_delete(store):
    first = store.create_chat("Alt")
    second = store.create_chat("Neu")
    store.add_exchange(first.id, "F", "A")                            # Alt wird zuletzt benutzt
    assert [c.title for c in store.chats()] == ["Alt", "Neu"]
    store.rename_chat(second.id, " Umbenannt ")
    assert store.get_chat(second.id).title == "Umbenannt"
    store.delete_chat(first.id)
    assert [c.title for c in store.chats()] == ["Umbenannt"]
    assert store.exchanges(first.id) == []


def test_projects_group_chats_and_deleting_keeps_chats(store):
    project = store.add_project("Bewerbungen")
    other = store.add_project("Alpha")
    chat = store.create_chat("Anschreiben")
    loose = store.create_chat("Sonstiges")
    store.set_chat_project(chat.id, project.id)
    assert [p.name for p in store.projects()] == ["Alpha", "Bewerbungen"]      # alphabetisch
    assert [c.id for c in store.chats_in_project(project.id)] == [chat.id]
    assert store.chats_in_project(other.id) == []
    store.rename_project(project.id, "Jobs")
    assert store.get_project(project.id).name == "Jobs"
    store.delete_project(project.id)
    assert store.get_project(project.id) is None
    assert store.get_chat(chat.id).project_id is None                # Chat bleibt erhalten
    assert store.get_chat(loose.id) is not None
    store.set_chat_project(chat.id, other.id)
    store.set_chat_project(chat.id, None)
    assert store.get_chat(chat.id).project_id is None


def test_chat_list_shows_only_chats_without_project(store):
    project = store.add_project("P")
    loose = store.create_chat("Lose")
    inside = store.create_chat("Im Projekt", project.id)
    assert [c.title for c in store.chats_without_project()] == ["Lose"]
    assert [c.title for c in store.chats_in_project(project.id)] == ["Im Projekt"]
    store.set_chat_project(inside.id, None)
    assert {c.title for c in store.chats_without_project()} == {"Lose", "Im Projekt"}
    assert store.chats_in_project(project.id) == []
    assert loose.id != inside.id


def test_copy_chat_duplicates_history_and_attachments_into_project(store):
    source_project = store.add_project("Quelle")
    target = store.add_project("Ziel")
    chat = store.create_chat("Original", source_project.id)
    store.add_exchange(chat.id, "F1", "A1")
    store.add_exchange(chat.id, "F2", "A2", "Kommentar", True, "Brief", "Brief")
    store.set_attachments(chat.id, [{"type": "file", "path": "C:\\a.txt"}])
    copy = store.copy_chat(chat.id, target.id)
    assert copy.id != chat.id and copy.project_id == target.id and copy.title == "Original"
    assert copy.attachments == [{"type": "file", "path": "C:\\a.txt"}]
    copied = store.exchanges(copy.id)
    assert [e.question for e in copied] == ["F2", "F1"] and copied[0].is_doc
    assert copied[0].word_filename == "Brief"
    assert len(store.exchanges(chat.id)) == 2                          # Original unverändert
    assert store.get_chat(chat.id).project_id == source_project.id
    store.delete_chat(chat.id)
    assert len(store.exchanges(copy.id)) == 2                          # Kopie ist unabhängig


def test_project_files_belong_to_one_project_and_go_with_it(store):
    a, b = store.add_project("A"), store.add_project("B")
    first = store.add_project_file(a.id, r"C:\x\Zebra.txt")
    store.add_project_file(a.id, r"C:\x\alpha.txt")
    store.add_project_file(b.id, r"C:\x\Zebra.txt")                  # dieselbe Datei in zwei Projekten
    assert store.add_project_file(a.id, r"C:\x\Zebra.txt").id == first.id     # doppelt: ein Eintrag
    assert [f.name for f in store.project_files(a.id)] == ["alpha.txt", "Zebra.txt"]
    assert [f.name for f in store.project_files(b.id)] == ["Zebra.txt"]
    assert store.get_project_file(first.id).project_id == a.id
    store.remove_project_file(first.id)
    assert [f.name for f in store.project_files(a.id)] == ["alpha.txt"]
    store.delete_project(a.id)                                       # Einträge gehen mit dem Projekt
    assert store.project_files(a.id) == [] and len(store.project_files(b.id)) == 1


def test_old_global_project_file_table_is_replaced(tmp_path):
    import sqlite3
    path = tmp_path / "alt.db"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE project_files (id INTEGER PRIMARY KEY, path TEXT NOT NULL UNIQUE, "
                 "added TEXT NOT NULL)")
    conn.execute("INSERT INTO project_files (path, added) VALUES ('C:\\alt.txt', '2026-09-19')")
    conn.commit()
    conn.close()
    s = Store(path)
    project = s.add_project("Neu")
    assert s.add_project_file(project.id, r"C:\neu.txt").project_id == project.id
    assert [f.name for f in s.project_files(project.id)] == ["neu.txt"]
    s.close()


def test_attachments_roundtrip(store):
    chat = store.create_chat("Mit Dateien")
    items = [{"type": "file", "path": "C:\\a\\b.txt"}, {"type": "folder", "path": "C:\\Ordner"}]
    store.set_attachments(chat.id, items)
    assert store.get_chat(chat.id).attachments == items


def test_title_from_question():
    assert title_from_question("  Wie   geht\ndas? ") == "Wie geht das?"
    assert len(title_from_question("x" * 200)) <= 62


def test_config_roundtrip_and_single_model_place(tmp_path):
    cfg = Config()
    cfg.thinking, cfg.chat_model = True, "qwen3:8b"
    cfg.save(tmp_path / "c.json")
    loaded = Config.load(tmp_path / "c.json")
    assert loaded.thinking is True and loaded.chat_model == "qwen3:8b"
    assert Config.load(tmp_path / "fehlt.json").chat_model == Config().chat_model


def test_unknown_backend_is_reported():
    with pytest.raises(LLMError):
        create_backend(Config(backend="gibt-es-nicht"))

# -- Antworttext in Zeilen -------------------------------------------------------------------------
from core import textlayout  # noqa: E402


def test_layout_puts_one_sentence_on_each_line_and_keeps_paragraphs():
    text = "Erster Satz. Zweiter Satz! Dritter Satz? Vierter Satz.\n\nNeuer Absatz. Noch ein Satz."
    assert textlayout.to_lines(text) == ("Erster Satz.\nZweiter Satz!\nDritter Satz?\nVierter Satz.\n\n"
                                         "Neuer Absatz.\nNoch ein Satz.")


@pytest.mark.parametrize("text", [
    "Das gilt z. B. für Sie. Danach geht es weiter.",
    "Er trifft Dr. Müller am 1. Mai in Freiburg.",
    "Es kostet ca. 5 Euro, vgl. Seite 3 und Nr. 7 usw. Weiter.",
    "Am 3.5. kommt er. Nein.",
    "Version 2.1 ist neu. Sie kostet 3,50 Euro.",
    "Wir treffen uns um 9.30 Uhr am Bahnhof.",
    "Siehe www.beispiel.de/seite. Danach unten.",
])
def test_layout_does_not_split_after_abbreviations_ordinals_and_numbers(text):
    lines = textlayout.to_lines(text).split("\n")
    for line in lines:                                   # keine Zeile beginnt mit Rest einer Abkürzung
        assert not line.startswith(("B.", "Müller", "Mai ", "Weiter", "Uhr", "5 "))
    assert "z. B. für Sie." in textlayout.to_lines("Das gilt z. B. für Sie. Danach geht es weiter.")


def test_layout_splits_after_years_and_amounts_but_not_before_lowercase():
    assert textlayout.to_lines("Das war im Jahr 2026. Danach kam es anders.").split("\n") == [
        "Das war im Jahr 2026.", "Danach kam es anders."]
    assert textlayout.to_lines("Er sagte: Ja. dann ging er.") == "Er sagte: Ja. dann ging er."


def test_layout_wraps_long_sentences_at_words_and_keeps_urls_whole():
    text = ("Dieser Satz ist sehr lang " * 8).strip() + ". Kurz. " + "https://" + "a" * 90 + ".example/x"
    lines = textlayout.to_lines(text).split("\n")
    assert all(len(line) <= textlayout.WIDTH for line in lines if "https://" not in line)
    assert any(line.startswith("https://" + "a" * 90) for line in lines)      # nie mitten im Wort
    assert "Kurz." in lines


def test_layout_keeps_lists_and_puts_sources_on_separate_lines():
    text = ("Das Wichtigste:\n- Erstens gilt das. Zweitens jenes.\n1. Punkt eins\n2) Punkt zwei\n\n\n\n"
            "Quellen aus dem Internet: https://a.example/1; https://b.example/2\n"
            "Gelesene Dateien: Vertrag.pdf, Seite 3")
    lines = textlayout.to_lines(text).split("\n")
    assert lines[:5] == ["Das Wichtigste:", "- Erstens gilt das. Zweitens jenes.", "1. Punkt eins",
                         "2) Punkt zwei", ""]
    assert lines[5:] == ["Quellen aus dem Internet:", "- https://a.example/1", "- https://b.example/2",
                         "Gelesene Dateien:", "- Vertrag.pdf, Seite 3"]


def test_layout_edge_cases():
    assert textlayout.to_lines("") == "" and textlayout.to_lines("\n\n  \n") == ""
    assert textlayout.to_lines("Zeile eins\u2028Zeile zwei") == "Zeile eins\nZeile zwei"
    assert textlayout.to_lines("Nur ein Satz ohne Punkt") == "Nur ein Satz ohne Punkt"
    assert textlayout.to_lines("Wirklich?! Ja... Nein") == "Wirklich?!\nJa...\nNein"
    assert textlayout.to_lines("Ende.\n\n") == "Ende."
